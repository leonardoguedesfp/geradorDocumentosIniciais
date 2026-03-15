"""Tests for placeholder_engine module."""

import os
import re
import pytest
from docx import Document
from docx.shared import Pt, RGBColor

from app.core.placeholder_engine import replace_placeholders, generate_document


@pytest.fixture
def simple_template(tmp_path):
    """Create a simple template with placeholders."""
    doc = Document()
    doc.add_paragraph("Nome: {{NOME_COMPLETO}}")
    doc.add_paragraph("CPF: {{CPF}}")
    doc.add_paragraph("Data: {{DATA_EXTENSO}}")
    path = str(tmp_path / "simple.docx")
    doc.save(path)
    return path


@pytest.fixture
def fragmented_template(tmp_path):
    """Create a template where a placeholder is split across multiple runs."""
    doc = Document()
    para = doc.add_paragraph("")
    # Simulate fragmented runs: {{NOME_ and COMPLETO}}
    run1 = para.add_run("Nome: {{NOME_")
    run2 = para.add_run("COMPLETO}}")
    # Make first run bold to test formatting preservation
    run1.bold = True
    path = str(tmp_path / "fragmented.docx")
    doc.save(path)
    return path


@pytest.fixture
def table_template(tmp_path):
    """Create a template with a table containing placeholders."""
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Nome"
    table.cell(0, 1).text = "{{NOME_COMPLETO}}"
    table.cell(1, 0).text = "CPF"
    table.cell(1, 1).text = "{{CPF}}"
    path = str(tmp_path / "table.docx")
    doc.save(path)
    return path


@pytest.fixture
def real_procuracao():
    """Path to the real Procuracao template."""
    path = os.path.join(os.path.dirname(__file__), "..", "templates_padrao", "Procuracao_TEMPLATE.docx")
    if os.path.exists(path):
        return path
    pytest.skip("Real template not available")


SAMPLE_PLACEHOLDERS = {
    "NOME_COMPLETO": "Maria da Silva Santos",
    "NACIONALIDADE": "brasileira",
    "ESTADO_CIVIL": "casada",
    "PROFISSAO": "professora",
    "RG": "1234567/SSP-DF",
    "CPF": "123.456.789-09",
    "ENDERECO_COMPLETO": "SQN 308, Bloco A, Apt 101 — Asa Norte, Brasília (DF), CEP 70747-010",
    "EMAIL": "maria@email.com",
    "TELEFONE": "(61) 99999-0000",
    "PARTE_CONTRARIA": "Empresa XYZ Ltda",
    "TIPO_ACAO": "Reclamação Trabalhista",
    "DATA_EXTENSO": "15 de março de 2026",
}


class TestReplacePlaceholders:
    def test_all_replaced(self, simple_template):
        doc = Document(simple_template)
        replace_placeholders(doc, SAMPLE_PLACEHOLDERS)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Maria da Silva Santos" in full_text
        assert "123.456.789-09" in full_text
        assert "{{" not in full_text

    def test_fragmented_runs(self, fragmented_template):
        doc = Document(fragmented_template)
        replace_placeholders(doc, SAMPLE_PLACEHOLDERS)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Maria da Silva Santos" in full_text
        assert "{{" not in full_text

    def test_table_placeholders(self, table_template):
        doc = Document(table_template)
        replace_placeholders(doc, SAMPLE_PLACEHOLDERS)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    assert "{{" not in cell.text

    def test_unknown_placeholder_replaced_empty(self, tmp_path):
        doc = Document()
        doc.add_paragraph("{{UNKNOWN_FIELD}}")
        path = str(tmp_path / "unknown.docx")
        doc.save(path)
        doc = Document(path)
        replace_placeholders(doc, SAMPLE_PLACEHOLDERS)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "{{" not in full_text

    def test_missing_data_replaced_empty(self, simple_template):
        doc = Document(simple_template)
        # Only provide partial data
        replace_placeholders(doc, {"NOME_COMPLETO": "Test"})
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "{{" not in full_text
        assert "Test" in full_text


class TestGenerateDocument:
    def test_generate_docx(self, simple_template, tmp_path):
        output = str(tmp_path / "output.docx")
        result = generate_document(simple_template, SAMPLE_PLACEHOLDERS, output)
        assert os.path.exists(result)
        # Verify content
        doc = Document(result)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Maria da Silva Santos" in full_text
        assert "{{" not in full_text

    def test_real_procuracao(self, real_procuracao, tmp_path):
        output = str(tmp_path / "procuracao_output.docx")
        result = generate_document(real_procuracao, SAMPLE_PLACEHOLDERS, output)
        assert os.path.exists(result)
        doc = Document(result)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Maria da Silva Santos" in full_text
        assert "{{" not in full_text


class TestFormattingPreservation:
    def test_bold_preserved(self, fragmented_template):
        doc = Document(fragmented_template)
        replace_placeholders(doc, SAMPLE_PLACEHOLDERS)
        # First run of the paragraph should be bold (from the original)
        para = doc.paragraphs[0]
        if para.runs:
            assert para.runs[0].bold is True
