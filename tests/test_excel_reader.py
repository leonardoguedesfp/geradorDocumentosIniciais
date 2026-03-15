"""Tests for excel_reader module."""

import os
import pytest
from openpyxl import Workbook
from app.core.excel_reader import read_excel, EXPECTED_COLUMNS


@pytest.fixture
def sample_excel(tmp_path):
    """Create a sample Excel file matching the expected structure."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"

    # Row 1: Title
    ws.cell(row=1, column=1, value="Ricardo Passos Advocacia — Cadastro de Clientes")
    # Row 2: Group headers
    ws.cell(row=2, column=1, value="DADOS PESSOAIS")
    # Row 3: Column headers
    headers = [
        "Nome Completo", "Nacionalidade", "Estado Civil", "Profissão", "RG", "CPF",
        "Endereço (Logradouro e Número)", "Bairro", "Cidade", "UF", "CEP",
        "E-mail", "Telefone(s)", "Parte Contrária / Réu", "Tipo de Ação",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(row=3, column=i, value=h)

    # Row 4: Client 1
    data1 = [
        "Maria da Silva Santos", "brasileira", "casada", "professora",
        "1234567/SSP-DF", "123.456.789-09", "SQN 308, Bloco A, Apt 101",
        "Asa Norte", "Brasília", "DF", "70747-010",
        "maria@email.com", "(61) 99999-0000", "Empresa XYZ Ltda", "Reclamação Trabalhista",
    ]
    for i, v in enumerate(data1, 1):
        ws.cell(row=4, column=i, value=v)

    # Row 5: Client 2
    data2 = [
        "João Pereira de Souza", "brasileiro", "solteiro", "engenheiro",
        "9876543/SSP-GO", "987.654.321-00", "Rua das Flores 123",
        "Centro", "Goiânia", "GO", "74000-000",
        "joao@email.com", "(62) 88888-0000", "Empresa ABC S.A.", "Ação Cível",
    ]
    for i, v in enumerate(data2, 1):
        ws.cell(row=5, column=i, value=v)

    path = str(tmp_path / "clientes.xlsx")
    wb.save(path)
    return path


@pytest.fixture
def excel_missing_fields(tmp_path):
    """Excel file with missing required fields in some rows."""
    wb = Workbook()
    ws = wb.active
    headers = [
        "Nome Completo", "Nacionalidade", "Estado Civil", "Profissão", "RG", "CPF",
        "Endereço (Logradouro e Número)", "Bairro", "Cidade", "UF", "CEP",
        "E-mail", "Telefone(s)", "Parte Contrária / Réu", "Tipo de Ação",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)

    # Row with missing name and CPF
    ws.cell(row=2, column=1, value="")
    ws.cell(row=2, column=2, value="brasileira")
    ws.cell(row=2, column=6, value="")
    ws.cell(row=2, column=7, value="Rua Teste 1")
    ws.cell(row=2, column=8, value="Bairro")
    ws.cell(row=2, column=9, value="Cidade")
    ws.cell(row=2, column=10, value="DF")
    ws.cell(row=2, column=11, value="70000-000")

    path = str(tmp_path / "missing.xlsx")
    wb.save(path)
    return path


@pytest.fixture
def excel_reordered_columns(tmp_path):
    """Excel file with columns in different order."""
    wb = Workbook()
    ws = wb.active
    # Reordered headers
    headers = [
        "CPF", "Nome Completo", "UF", "Cidade", "Nacionalidade",
        "Estado Civil", "Profissão", "RG",
        "Endereço (Logradouro e Número)", "Bairro", "CEP",
        "E-mail", "Telefone(s)", "Parte Contrária / Réu", "Tipo de Ação",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)

    # Data
    ws.cell(row=2, column=1, value="111.222.333-44")
    ws.cell(row=2, column=2, value="Test Client")
    ws.cell(row=2, column=3, value="DF")
    ws.cell(row=2, column=4, value="Brasília")
    ws.cell(row=2, column=5, value="brasileira")
    ws.cell(row=2, column=6, value="solteiro")
    ws.cell(row=2, column=7, value="advogado")
    ws.cell(row=2, column=8, value="555555")
    ws.cell(row=2, column=9, value="Rua A 10")
    ws.cell(row=2, column=10, value="Centro")
    ws.cell(row=2, column=11, value="70000-000")
    ws.cell(row=2, column=12, value="test@mail.com")
    ws.cell(row=2, column=13, value="(61) 99999-0000")
    ws.cell(row=2, column=14, value="Réu Teste")
    ws.cell(row=2, column=15, value="Ação Teste")

    path = str(tmp_path / "reordered.xlsx")
    wb.save(path)
    return path


@pytest.fixture
def real_excel():
    """Path to the real Excel template."""
    path = os.path.join(os.path.dirname(__file__), "..", "templates_padrao", "Cadastro_Clientes_RicardoPassos.xlsx")
    if os.path.exists(path):
        return path
    pytest.skip("Real Excel template not available")


class TestReadExcel:
    def test_read_two_clients(self, sample_excel):
        clientes, warnings, row_warnings = read_excel(sample_excel)
        assert len(clientes) == 2
        assert clientes[0].nome_completo == "Maria da Silva Santos"
        assert clientes[1].nome_completo == "João Pereira de Souza"
        assert len(warnings) == 0

    def test_missing_required_fields(self, excel_missing_fields):
        clientes, warnings, row_warnings = read_excel(excel_missing_fields)
        assert len(clientes) == 1
        assert len(row_warnings) > 0

    def test_reordered_columns(self, excel_reordered_columns):
        clientes, warnings, row_warnings = read_excel(excel_reordered_columns)
        assert len(clientes) == 1
        assert clientes[0].nome_completo == "Test Client"
        assert clientes[0].cpf == "111.222.333-44"
        assert clientes[0].uf == "DF"

    def test_real_excel_structure(self, real_excel):
        clientes, warnings, row_warnings = read_excel(real_excel)
        # The real template is empty (headers only), so no clients
        assert isinstance(clientes, list)
        assert len(warnings) == 0  # Headers should be found

    def test_multiple_clients_all_fields(self, sample_excel):
        clientes, _, _ = read_excel(sample_excel)
        for c in clientes:
            assert c.nacionalidade != ""
            assert c.estado_civil != ""
            assert c.cpf != ""
            assert c.logradouro_numero != ""


class TestBatchProcessing:
    def test_all_clients_returned(self, sample_excel):
        clientes, _, _ = read_excel(sample_excel)
        assert len(clientes) == 2

    def test_individual_failure_doesnt_block(self, excel_missing_fields):
        """Even rows with missing data are returned (with warnings)."""
        clientes, _, row_warnings = read_excel(excel_missing_fields)
        assert len(clientes) >= 1
