"""Tests for template_loader module."""

import os
import pytest
from pathlib import Path
from docx import Document

from app.core.template_loader import (
    check_templates_folder,
    extract_placeholders,
    validate_template,
    load_template,
    EXPECTED_FILES,
    EXPECTED_PLACEHOLDERS,
)


@pytest.fixture
def templates_dir():
    """Return path to the real templates_padrao directory."""
    base = Path(__file__).resolve().parent.parent / "templates_padrao"
    if base.exists():
        return str(base)
    pytest.skip("templates_padrao directory not found")


@pytest.fixture
def temp_templates(tmp_path):
    """Create temporary template files for testing."""
    for doc_type, filename in EXPECTED_FILES.items():
        doc = Document()
        doc.add_paragraph("Test {{NOME_COMPLETO}} and {{CPF}}")
        doc.save(str(tmp_path / filename))
    return str(tmp_path)


@pytest.fixture
def corrupt_template(tmp_path):
    """Create a corrupt (non-docx) file."""
    path = tmp_path / "Procuracao_TEMPLATE.docx"
    path.write_text("this is not a docx file")
    return str(path)


class TestCheckTemplatesFolder:
    def test_all_present(self, templates_dir):
        result = check_templates_folder(templates_dir)
        for doc_type in EXPECTED_FILES:
            assert result[doc_type] is not None

    def test_missing_file(self, tmp_path):
        # Only create one file
        doc = Document()
        doc.save(str(tmp_path / "Procuracao_TEMPLATE.docx"))
        result = check_templates_folder(str(tmp_path))
        assert result["procuracao"] is not None
        assert result["declaracao"] is None
        assert result["contrato"] is None

    def test_empty_folder(self, tmp_path):
        result = check_templates_folder(str(tmp_path))
        for doc_type in EXPECTED_FILES:
            assert result[doc_type] is None


class TestExtractPlaceholders:
    def test_real_procuracao(self, templates_dir):
        path = os.path.join(templates_dir, "Procuracao_TEMPLATE.docx")
        placeholders = extract_placeholders(path)
        assert "NOME_COMPLETO" in placeholders
        assert "CPF" in placeholders
        assert "DATA_EXTENSO" in placeholders

    def test_real_declaracao(self, templates_dir):
        path = os.path.join(templates_dir, "Declaracao_Hipossuficiencia_TEMPLATE.docx")
        placeholders = extract_placeholders(path)
        assert "TIPO_ACAO" in placeholders
        assert "PARTE_CONTRARIA" in placeholders

    def test_real_contrato(self, templates_dir):
        path = os.path.join(templates_dir, "Contrato_Prestacao_Servicos_TEMPLATE.docx")
        placeholders = extract_placeholders(path)
        assert "TIPO_ACAO" in placeholders
        assert "NOME_COMPLETO" in placeholders

    def test_corrupt_file(self, corrupt_template):
        with pytest.raises(ValueError):
            extract_placeholders(corrupt_template)


class TestValidateTemplate:
    def test_valid_procuracao(self, templates_dir):
        path = os.path.join(templates_dir, "Procuracao_TEMPLATE.docx")
        is_valid, warnings = validate_template(path, "procuracao")
        assert is_valid is True

    def test_valid_declaracao(self, templates_dir):
        path = os.path.join(templates_dir, "Declaracao_Hipossuficiencia_TEMPLATE.docx")
        is_valid, warnings = validate_template(path, "declaracao")
        assert is_valid is True

    def test_valid_contrato(self, templates_dir):
        path = os.path.join(templates_dir, "Contrato_Prestacao_Servicos_TEMPLATE.docx")
        is_valid, warnings = validate_template(path, "contrato")
        assert is_valid is True

    def test_corrupt_template(self, corrupt_template):
        is_valid, warnings = validate_template(corrupt_template, "procuracao")
        assert is_valid is False
        assert len(warnings) > 0

    def test_missing_placeholders_warning(self, tmp_path):
        # Template with only one placeholder
        doc = Document()
        doc.add_paragraph("{{NOME_COMPLETO}}")
        path = str(tmp_path / "test.docx")
        doc.save(path)
        is_valid, warnings = validate_template(path, "procuracao")
        assert is_valid is True
        assert len(warnings) > 0  # Missing placeholders warning


class TestLoadTemplate:
    def test_load_valid(self, templates_dir):
        path = os.path.join(templates_dir, "Procuracao_TEMPLATE.docx")
        doc = load_template(path)
        assert doc is not None
        assert len(doc.paragraphs) > 0

    def test_load_corrupt(self, corrupt_template):
        with pytest.raises(ValueError):
            load_template(corrupt_template)
