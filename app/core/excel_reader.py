"""Read client data from Excel spreadsheet (Microsoft Forms export)."""

from pathlib import Path
from openpyxl import load_workbook
from app.models.cliente import Cliente


# Expected column headers (case-insensitive matching)
EXPECTED_COLUMNS = [
    "Nome Completo",
    "Nacionalidade",
    "Estado Civil",
    "Profissão",
    "RG",
    "CPF",
    "Endereço (Logradouro e Número)",
    "Bairro",
    "Cidade",
    "UF",
    "CEP",
    "E-mail",
    "Telefone(s)",
    "Parte Contrária / Réu",
    "Tipo de Ação",
]

COLUMN_MAPPING = {
    "nome completo": "nome_completo",
    "nacionalidade": "nacionalidade",
    "estado civil": "estado_civil",
    "profissão": "profissao",
    "profissao": "profissao",
    "rg": "rg",
    "cpf": "cpf",
    "endereço (logradouro e número)": "logradouro_numero",
    "endereco (logradouro e numero)": "logradouro_numero",
    "bairro": "bairro",
    "cidade": "cidade",
    "uf": "uf",
    "cep": "cep",
    "e-mail": "email",
    "email": "email",
    "telefone(s)": "telefone",
    "telefones": "telefone",
    "telefone": "telefone",
    "parte contrária / réu": "parte_contraria",
    "parte contraria / reu": "parte_contraria",
    "tipo de ação": "tipo_acao",
    "tipo de acao": "tipo_acao",
}

REQUIRED_FIELDS = {
    "nome_completo", "nacionalidade", "estado_civil", "rg", "cpf",
    "logradouro_numero", "bairro", "cidade", "uf", "cep",
    "parte_contraria", "tipo_acao",
}


def _find_header_row(ws) -> tuple[int, dict[int, str]]:
    """Find the header row and return (row_number, {col_index: field_name})."""
    for row_idx in range(1, min(ws.max_row + 1, 20)):
        col_map = {}
        for col_idx in range(1, ws.max_column + 1):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is None:
                continue
            normalized = str(val).strip().lower()
            if normalized in COLUMN_MAPPING:
                col_map[col_idx] = COLUMN_MAPPING[normalized]
        # Consider it a header row if we found at least 5 known columns
        if len(col_map) >= 5:
            return row_idx, col_map
    return -1, {}


def read_excel(file_path: str) -> tuple[list[Cliente], list[str], list[tuple[int, list[str]]]]:
    """Read clients from an Excel file.

    Returns:
        - List of Cliente objects
        - List of missing column warnings
        - List of (row_number, [missing_field_names]) for rows with missing required data
    """
    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active

    header_row, col_map = _find_header_row(ws)
    if header_row < 0:
        return [], ["Não foi possível encontrar a linha de cabeçalho na planilha."], []

    # Check for missing expected columns
    found_fields = set(col_map.values())
    missing_columns = []
    for expected in EXPECTED_COLUMNS:
        normalized = expected.strip().lower()
        if normalized in COLUMN_MAPPING:
            field = COLUMN_MAPPING[normalized]
            if field not in found_fields:
                missing_columns.append(expected)

    clientes = []
    row_warnings = []

    for row_idx in range(header_row + 1, ws.max_row + 1):
        # Check if row has any data
        values = {}
        has_data = False
        for col_idx, field_name in col_map.items():
            cell_val = ws.cell(row=row_idx, column=col_idx).value
            if cell_val is not None:
                has_data = True
                values[field_name] = str(cell_val).strip()
            else:
                values[field_name] = ""

        if not has_data:
            continue

        # Check required fields
        missing_fields = []
        for req in REQUIRED_FIELDS:
            if req in found_fields and not values.get(req, ""):
                missing_fields.append(req)

        if missing_fields:
            row_warnings.append((row_idx, missing_fields))

        cliente = Cliente(
            nome_completo=values.get("nome_completo", ""),
            nacionalidade=values.get("nacionalidade", ""),
            estado_civil=values.get("estado_civil", ""),
            profissao=values.get("profissao", ""),
            rg=values.get("rg", ""),
            cpf=values.get("cpf", ""),
            logradouro_numero=values.get("logradouro_numero", ""),
            bairro=values.get("bairro", ""),
            cidade=values.get("cidade", ""),
            uf=values.get("uf", ""),
            cep=values.get("cep", ""),
            email=values.get("email", ""),
            telefone=values.get("telefone", ""),
            parte_contraria=values.get("parte_contraria", ""),
            tipo_acao=values.get("tipo_acao", ""),
        )
        clientes.append(cliente)

    wb.close()
    return clientes, missing_columns, row_warnings
