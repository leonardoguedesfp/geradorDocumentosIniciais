"""Engine to substitute {{PLACEHOLDER}} values in .docx documents.

Handles fragmented runs where a single placeholder may span multiple runs.
"""

import copy
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def _replace_in_paragraph(paragraph, placeholders: dict[str, str]) -> None:
    """Replace placeholders in a paragraph, handling fragmented runs."""
    runs = paragraph.runs
    if not runs:
        return

    # Concatenate all run texts
    full_text = "".join(run.text for run in runs)

    # Check if any placeholder is present
    if "{{" not in full_text:
        return

    # Perform all substitutions
    new_text = full_text
    for key, value in placeholders.items():
        new_text = new_text.replace("{{" + key + "}}", value)

    # Also replace any remaining {{...}} with empty string
    new_text = re.sub(r"\{\{\w+\}\}", "", new_text)

    if new_text == full_text:
        return

    # Rebuild runs: put all text in first run, clear the rest
    # Preserve formatting of the first run
    if runs:
        runs[0].text = new_text
        for run in runs[1:]:
            run.text = ""


def _replace_in_table(table, placeholders: dict[str, str]) -> None:
    """Replace placeholders in all cells of a table."""
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                _replace_in_paragraph(paragraph, placeholders)


def replace_placeholders(doc: Document, placeholders: dict[str, str]) -> Document:
    """Replace all {{PLACEHOLDER}} in the document with provided values.

    Args:
        doc: A python-docx Document object.
        placeholders: Dict mapping placeholder names (without braces) to values.

    Returns:
        The modified Document object.
    """
    for paragraph in doc.paragraphs:
        _replace_in_paragraph(paragraph, placeholders)

    for table in doc.tables:
        _replace_in_table(table, placeholders)

    return doc


def generate_document(
    template_path: str,
    placeholders: dict[str, str],
    output_path: str,
) -> str:
    """Load a template, replace placeholders, and save to output_path.

    Returns the output path on success.
    """
    doc = Document(template_path)
    replace_placeholders(doc, placeholders)
    doc.save(output_path)
    return output_path
