"""Load and validate .docx templates."""

import re
from pathlib import Path
from docx import Document


EXPECTED_FILES = {
    "procuracao": "Procuracao_TEMPLATE.docx",
    "declaracao": "Declaracao_Hipossuficiencia_TEMPLATE.docx",
    "contrato": "Contrato_Prestacao_Servicos_TEMPLATE.docx",
}

EXPECTED_PLACEHOLDERS = {
    "procuracao": {
        "NOME_COMPLETO", "NACIONALIDADE", "ESTADO_CIVIL", "PROFISSAO",
        "RG", "CPF", "ENDERECO_COMPLETO", "EMAIL", "TELEFONE",
        "PARTE_CONTRARIA", "DATA_EXTENSO",
    },
    "declaracao": {
        "NOME_COMPLETO", "NACIONALIDADE", "ESTADO_CIVIL",
        "RG", "CPF", "ENDERECO_COMPLETO",
        "PARTE_CONTRARIA", "TIPO_ACAO", "DATA_EXTENSO",
    },
    "contrato": {
        "NOME_COMPLETO", "NACIONALIDADE", "ESTADO_CIVIL",
        "RG", "CPF", "ENDERECO_COMPLETO",
        "PARTE_CONTRARIA", "TIPO_ACAO", "DATA_EXTENSO",
    },
}


def check_templates_folder(folder: str) -> dict[str, str | None]:
    """Check which template files are present in the folder.

    Returns dict mapping doc_type -> full path (or None if missing).
    """
    result = {}
    folder_path = Path(folder)
    for doc_type, filename in EXPECTED_FILES.items():
        path = folder_path / filename
        result[doc_type] = str(path) if path.exists() else None
    return result


def extract_placeholders(doc_path: str) -> set[str]:
    """Extract all {{PLACEHOLDER}} names from a .docx file."""
    try:
        doc = Document(doc_path)
    except Exception:
        raise ValueError(f"Não foi possível abrir o template: {doc_path}")

    placeholders = set()
    for para in doc.paragraphs:
        text = "".join(run.text for run in para.runs)
        found = re.findall(r"\{\{(\w+)\}\}", text)
        placeholders.update(found)
        # Also check paragraph text directly (for non-run text)
        found2 = re.findall(r"\{\{(\w+)\}\}", para.text)
        placeholders.update(found2)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    text = "".join(run.text for run in para.runs)
                    found = re.findall(r"\{\{(\w+)\}\}", text)
                    placeholders.update(found)

    return placeholders


def validate_template(doc_path: str, doc_type: str) -> tuple[bool, list[str]]:
    """Validate a template file. Returns (is_valid, list_of_warnings).

    Warnings include missing expected placeholders (non-blocking).
    """
    warnings = []
    try:
        found = extract_placeholders(doc_path)
    except ValueError as e:
        return False, [str(e)]

    expected = EXPECTED_PLACEHOLDERS.get(doc_type, set())
    missing = expected - found
    if missing:
        warnings.append(
            f"Placeholders esperados ausentes: {', '.join(sorted(missing))}"
        )

    return True, warnings


def load_template(doc_path: str) -> Document:
    """Load a .docx Document object from path."""
    try:
        return Document(doc_path)
    except Exception:
        raise ValueError(f"Template corrompido ou inválido: {doc_path}")
